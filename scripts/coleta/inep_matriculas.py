#!/usr/bin/env python3
"""Matrículas da educação básica por UF e dependência administrativa (Inep, Sinopse Estatística).

Lê a tabela 1.2 (localização e dependência administrativa) da Sinopse Estatística da
Educação Básica de cada ano e grava, por UF: total, federal, estadual, municipal e privada.
Fonte: https://download.inep.gov.br/dados_abertos/sinopses_estatisticas/sinopse_estatistica_censo_escolar_{ano}.zip
Saída: data/aux/coleta/inep_matriculas/matriculas_uf.csv
"""
import sys, zipfile, pathlib, csv, re, unicodedata
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from _rede import baixar
import openpyxl
OUT = pathlib.Path(__file__).resolve().parents[2] / "data" / "aux" / "coleta" / "inep_matriculas"
OUT.mkdir(parents=True, exist_ok=True)
UF = {"Rondônia": "RO", "Acre": "AC", "Amazonas": "AM", "Roraima": "RR", "Pará": "PA", "Amapá": "AP", "Tocantins": "TO",
      "Maranhão": "MA", "Piauí": "PI", "Ceará": "CE", "Rio Grande do Norte": "RN", "Paraíba": "PB", "Pernambuco": "PE",
      "Alagoas": "AL", "Sergipe": "SE", "Bahia": "BA", "Minas Gerais": "MG", "Espírito Santo": "ES", "Rio de Janeiro": "RJ",
      "São Paulo": "SP", "Paraná": "PR", "Santa Catarina": "SC", "Rio Grande do Sul": "RS", "Mato Grosso do Sul": "MS",
      "Mato Grosso": "MT", "Goiás": "GO", "Distrito Federal": "DF"}
norm = lambda s: unicodedata.normalize("NFC", str(s or "")).strip()
num = lambda v: int(float(v)) if v not in (None, "", "-", "--") else 0
saida = []
for ano in [int(a) for a in sys.argv[1:]] or list(range(2017, 2026)):
    z = f"/tmp/sinopse_{ano}.zip"
    nomes_url = [f"https://download.inep.gov.br/dados_abertos/sinopses_estatisticas/sinopse_estatistica_censo_escolar_{ano}.zip",
                 f"https://download.inep.gov.br/dados_abertos/sinopses_estatisticas/sinopses_estatisticas_censo_escolar_{ano}.zip",
                 f"https://download.inep.gov.br/informacoes_estatisticas/sinopses_estatisticas/sinopses_educacao_basica/sinopse_estatistica_da_educacao_basica_{ano}.zip",
                 f"https://download.inep.gov.br/informacoes_estatisticas/sinopses_estatisticas/sinopses_educacao_basica/sinopse_estatistica_educacao_basica_{ano}.zip"]
    ok = False
    for u in nomes_url:
        try: baixar(u, z, tentativas=3); ok = True; print(ano, "baixado de", u, flush=True); break
        except Exception as e: print(ano, "falhou", u, e, flush=True)
    if not ok: continue
    with zipfile.ZipFile(z) as zf:
        xl = [n for n in zf.namelist() if n.lower().endswith(".xlsx")]
        if not xl: print(ano, "sem xlsx", zf.namelist()); continue
        p = zf.extract(xl[0], f"/tmp/sin{ano}")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    alvo = None
    for aba in wb.sheetnames:
        if "sum" in aba.lower() or "resumo" in aba.lower(): continue
        ws = wb[aba]
        linhas_topo = [" ".join(norm(c) for c in r if c) for r in ws.iter_rows(min_row=1, max_row=10, values_only=True)]
        titulo = [t for t in linhas_topo if re.match(r"^1\.2\s*[–-]", t)]
        if titulo and "Dependência Administrativa" in titulo[0] and "Localização" in titulo[0] and "Educação Básica" in titulo[0]:
            alvo = ws; print(ano, "aba", aba, "|", titulo[0][:120], flush=True); break
    if alvo is None: print(ano, "tabela 1.2 não encontrada"); continue
    n = 0
    for r in alvo.iter_rows(values_only=True):
        uf, mun = norm(r[1]) if len(r) > 1 else "", norm(r[2]) if len(r) > 2 else ""
        if uf in UF and not mun:
            v = [num(x) for x in r[4:15]]
            tot, ut, uF, uE, uM, uP, rt, rF, rE, rM, rP = v
            assert ut == uF + uE + uM + uP and rt == rF + rE + rM + rP and tot == ut + rt, (ano, uf, v)
            saida.append({"ano": ano, "uf": UF[uf], "total": tot, "federal": uF + rF, "estadual": uE + rE,
                          "municipal": uM + rM, "privada": uP + rP})
            n += 1
    print(ano, n, "UFs", flush=True)
with open(OUT / "matriculas_uf.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(saida[0].keys()), delimiter=";"); w.writeheader(); w.writerows(saida)
print("fim")
