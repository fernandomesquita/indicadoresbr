#!/usr/bin/env python3
"""Baixa a Sinopse Estatística do Censo Escolar (Inep) e registra a estrutura das planilhas."""
import sys, zipfile, pathlib, json
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from _rede import baixar
import openpyxl
OUT = pathlib.Path(__file__).resolve().parents[2] / "data" / "aux" / "coleta" / "inep_sinopse"
OUT.mkdir(parents=True, exist_ok=True)
anos = [int(a) for a in sys.argv[1:]] or [2024]
for ano in anos:
    url = f"https://download.inep.gov.br/dados_abertos/sinopses_estatisticas/sinopse_estatistica_censo_escolar_{ano}.zip"
    z = f"/tmp/sinopse_{ano}.zip"
    baixar(url, z)
    with zipfile.ZipFile(z) as zf:
        nomes = zf.namelist(); print(ano, nomes, flush=True)
        (OUT / f"arquivos_{ano}.json").write_text(json.dumps(nomes, ensure_ascii=False, indent=1))
        for n in nomes:
            if not n.lower().endswith((".xlsx", ".xls")): continue
            zf.extract(n, "/tmp/sin"); p = pathlib.Path("/tmp/sin") / n
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            linhas = [f"ARQUIVO {n}", "ABAS: " + " | ".join(wb.sheetnames)]
            for aba in wb.sheetnames[:12]:
                ws = wb[aba]; linhas.append(f"\n=== {aba}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 25: break
                    linhas.append(" ; ".join("" if v is None else str(v) for v in row[:30]))
            (OUT / f"estrutura_{ano}_{pathlib.Path(n).stem[:40]}.txt").write_text("\n".join(linhas), encoding="utf-8")
print("fim")
